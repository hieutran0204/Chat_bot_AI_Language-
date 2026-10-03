# name: recursive_chunker.py
# description: IChunker adapter using LangChain's RecursiveCharacterTextSplitter.
#              Supports PDF (PyMuPDF), DOCX (python-docx), and plain TXT parsing.

import logging
from pathlib import Path

from app.core.interfaces.chunker import IChunker, TextChunk

logger = logging.getLogger(__name__)


class ChunkerError(Exception):
    """Raised when a document cannot be parsed or chunked."""


class RecursiveChunker(IChunker):
    """
    Document parser and text splitter using RecursiveCharacterTextSplitter.

    Prioritises semantic boundaries (paragraph → sentence → word) to avoid
    cutting mid-sentence. Supported formats: PDF, DOCX, TXT.

    Args:
        chunk_size: Maximum characters per chunk (default: 512).
        chunk_overlap: Character overlap between adjacent chunks (default: 64).
    """

    def __init__(
        self,
        chunk_size: int = 512,
        chunk_overlap: int = 64,
    ) -> None:
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def _get_splitter(self):
        """Build a LangChain splitter that prioritises paragraph/sentence breaks."""
        from langchain_text_splitters import RecursiveCharacterTextSplitter

        return RecursiveCharacterTextSplitter(
            chunk_size=self.chunk_size,
            chunk_overlap=self.chunk_overlap,
            separators=["\n\n", "\n", ". ", "! ", "? ", " ", ""],
        )

    # ── Parsers ───────────────────────────────────────────────────────────────

    def _parse_pdf(self, file_path: Path) -> list[tuple[str, dict]]:
        """
        Extract text from a PDF file, page by page.

        Logs a warning when no text is extracted from any page, which typically
        indicates a scanned/image-only PDF that requires OCR pre-processing.

        Args:
            file_path: Absolute path to the PDF file.

        Returns:
            List of (text, metadata) tuples, one per non-empty page.

        Raises:
            ChunkerError: If the PDF cannot be opened or read.
        """
        import fitz  # PyMuPDF

        try:
            pages: list[tuple[str, dict]] = []
            with fitz.open(str(file_path)) as doc:
                for page_num, page in enumerate(doc, start=1):
                    text = page.get_text("text").strip()
                    if text:
                        pages.append((text, {"page": page_num, "filename": file_path.name}))

            if not pages:
                logger.warning(
                    "No text extracted from PDF '%s' — the file may be a scanned image "
                    "without a text layer (OCR required) or may be encrypted.",
                    file_path.name,
                )
            else:
                logger.info("Parsed %d pages from PDF: %s", len(pages), file_path.name)

            return pages
        except Exception as exc:
            raise ChunkerError(f"Failed to parse PDF '{file_path.name}': {exc}") from exc

    def _parse_docx(self, file_path: Path) -> list[tuple[str, dict]]:
        """
        Extract text from a DOCX file, grouped as a single block.

        Includes content from both paragraphs and tables. python-docx's
        doc.paragraphs does not cover text inside tables, so table cells are
        iterated separately and appended as pipe-delimited rows.

        Args:
            file_path: Absolute path to the DOCX file.

        Returns:
            Single-element list with full document text and metadata.

        Raises:
            ChunkerError: If the DOCX cannot be opened or read.
        """
        try:
            from docx import Document

            doc = Document(str(file_path))

            # --- Paragraph text ---
            paragraph_text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

            # --- Table text (omitted by doc.paragraphs) ---
            # Each row is serialised as "cell1 | cell2 | ..." so that structured
            # data (contracts, reports) survives chunking in a readable form.
            table_rows: list[str] = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        table_rows.append(row_text)

            parts = [paragraph_text] if paragraph_text else []
            if table_rows:
                parts.append("\n".join(table_rows))

            full_text = "\n\n".join(parts)
            logger.info(
                "Parsed DOCX: %s (%d table rows extracted)",
                file_path.name,
                len(table_rows),
            )
            return [(full_text, {"filename": file_path.name})]

        except Exception as exc:
            raise ChunkerError(f"Failed to parse DOCX '{file_path.name}': {exc}") from exc

    def _parse_txt(self, file_path: Path) -> list[tuple[str, dict]]:
        """
        Read a plain text file.

        Args:
            file_path: Absolute path to the TXT file.

        Returns:
            Single-element list with file content and metadata.

        Raises:
            ChunkerError: If the file cannot be read.
        """
        try:
            text = file_path.read_text(encoding="utf-8", errors="replace")
            logger.info("Parsed TXT: %s", file_path.name)
            return [(text, {"filename": file_path.name})]
        except Exception as exc:
            raise ChunkerError(f"Failed to read TXT '{file_path.name}': {exc}") from exc

    # ── Public API ────────────────────────────────────────────────────────────

    def chunk_file(self, file_path: str | Path) -> list[TextChunk]:
        """
        Parse a document and split it into overlapping text chunks.

        Dispatches to the correct parser based on file extension,
        then applies the recursive character splitter.

        Args:
            file_path: Path to the document file (PDF, DOCX, or TXT).

        Returns:
            Ordered list of TextChunk objects ready for embedding.

        Raises:
            FileNotFoundError: If the file does not exist at the given path.
            ValueError: If the file extension is not supported.
            ChunkerError: If the file cannot be parsed (corrupted, encrypted, etc.).
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Document not found: {path}")

        ext = path.suffix.lower().lstrip(".")
        parsers = {
            "pdf": self._parse_pdf,
            "docx": self._parse_docx,
            "txt": self._parse_txt,
        }

        if ext not in parsers:
            raise ValueError(f"Unsupported file type: .{ext}. Allowed: {list(parsers)}")

        raw_sections = parsers[ext](path)
        return self._split_sections(raw_sections)

    def chunk_text(self, text: str, metadata: dict | None = None) -> list[TextChunk]:
        """
        Split a raw text string into chunks without file I/O.

        Args:
            text: The raw input text.
            metadata: Optional metadata dict to attach to every chunk.

        Returns:
            Ordered list of TextChunk objects.
        """
        return self._split_sections([(text, metadata or {})])

    def _split_sections(self, sections: list[tuple[str, dict]]) -> list[TextChunk]:
        """
        Apply the recursive splitter to a list of (text, metadata) sections.

        Args:
            sections: List of (text, metadata) tuples from a parser.

        Returns:
            Flat ordered list of TextChunk objects.
        """
        splitter = self._get_splitter()
        chunks: list[TextChunk] = []

        for text, meta in sections:
            for split_text in splitter.split_text(text):
                clean = split_text.strip()
                if len(clean) < 20:  # Skip noise / header fragments
                    continue
                chunks.append(
                    TextChunk(
                        content=clean,
                        chunk_index=len(chunks),
                        metadata=meta,
                    )
                )

        logger.info(
            "RecursiveChunker: produced %d chunks (size=%d, overlap=%d)",
            len(chunks),
            self.chunk_size,
            self.chunk_overlap,
        )
        return chunks
